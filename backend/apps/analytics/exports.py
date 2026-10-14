"""CSV and XLSX writers for the analytics portal.

CSV needs nothing. XLSX needs ``openpyxl``, which is a real dependency with a
real failure mode: a deployment that upgrades the code before the requirements
lock is rebuilt will have no ``openpyxl``, and an export that raises
``ImportError`` at request time is a 500 with a stack trace and no explanation.
So the import is probed and its absence is a **503 naming the dependency**.
CSV keeps working in that deployment — the point of degrading.

Three rules the rest of the file exists to enforce:

**Same data as the screen.** :func:`build_report` calls the same builder the
on-screen endpoint calls, with the same query params, and takes the ``table``
that builder returns. There is no second query path to drift, and rendering is
a separate step so a caller can build once and render once.

**No silent truncation.** A row cap that quietly drops the tail produces a file
that lies. Over the cap is a 400 naming the row count and the cap.

**Cells cannot execute.** A trending-post excerpt is user text. A cell starting
with ``=``, ``+``, ``-`` or ``@`` is a formula in Excel and Sheets, so string
cells are prefixed with an apostrophe, which those tools render as text.
Numbers are left alone — prefixing them would turn a total into a label.
"""
import csv
import io
import logging
from dataclasses import dataclass

from django.http import HttpResponse
from django.utils import timezone
from rest_framework.exceptions import ValidationError

logger = logging.getLogger(__name__)

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
CSV_MIME = 'text/csv'
FORMATS = ('csv', 'xlsx')

#: Hard cap on exported rows. 50k x ~10 columns is a spreadsheet a human can
#: still open; past that the answer is "narrow the range", not a bigger file.
DEFAULT_MAX_ROWS = 50000

#: Characters Excel and Google Sheets treat as the start of a formula.
_FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def max_rows():
    from django.conf import settings
    return int(getattr(settings, 'ANALYTICS_EXPORT_MAX_ROWS', DEFAULT_MAX_ROWS))


def openpyxl_available():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def _cell(value):
    """One CSV/XLSX-safe cell. Never raises, never returns None."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text[:1] in _FORMULA_PREFIXES:
        return "'" + text
    return text


def _row(values):
    return [_cell(v) for v in values]


def enforce_row_cap(built):
    """Raise rather than truncate. A short file is a lie about the data."""
    rows = built.table.get('rows') or []
    limit = max_rows()
    if len(rows) > limit:
        raise ValidationError({
            'from': [
                f'The {built.key} export would contain {len(rows):,} rows, above the '
                f'{limit:,} row cap. Narrow the date range (a shorter window '
                'produces fewer buckets) and try again. The cap is '
                'ANALYTICS_EXPORT_MAX_ROWS. The file is not truncated silently '
                'because a truncated export would differ from the report you '
                'were looking at.'
            ],
        })
    return rows


@dataclass
class BuiltReport:
    """A report that has been computed once and can now be rendered."""

    key: str
    label: str
    table: dict
    payload: dict
    params: dict


def build_report(report_key, params):
    """Run one report builder and return a :class:`BuiltReport`.

    Raises DRF ``ValidationError`` for an unknown report or an invalid window —
    identical messages to the on-screen endpoint, which is the whole point.
    """
    from .platform_metrics import REPORTS

    report = REPORTS.get(report_key)
    if report is None:
        raise ValidationError({
            'report': [
                f'Unknown report {report_key!r}. Must be one of: '
                f'{", ".join(sorted(REPORTS))}.',
            ],
        })
    payload = report['builder'](params)
    table = payload.get('table')
    if not table:
        # A builder that forgot its table would otherwise export an empty file
        # that looks like "no data". Fail loudly instead.
        raise ValidationError({
            'report': [f'The {report_key!r} report produced no table to export.'],
        })
    return BuiltReport(
        key=report_key, label=report['label'], table=table, payload=payload,
        params={k: v for k, v in dict(params).items()
                if k not in ('format',)},
    )


def _filename(built, extension):
    window = built.payload.get('window') or {}
    start = window.get('from') or 'start'
    end = window.get('to') or 'end'
    return f'analytics-{built.key}_{start}_{end}.{extension}'


def _attach(response, built, extension, row_count):
    window = built.payload.get('window') or {}
    response['Content-Disposition'] = f'attachment; filename="{_filename(built, extension)}"'
    # Provenance in the headers too: a downloaded file keeps the filename but
    # loses these, which is why the xlsx writer also writes a provenance sheet.
    response['X-Analytics-Report'] = built.key
    response['X-Analytics-From'] = str(window.get('from') or '')
    response['X-Analytics-To'] = str(window.get('to') or '')
    response['X-Analytics-Rows'] = str(row_count)
    data_source = built.payload.get('data_source')
    if data_source:
        response['X-Analytics-Data-Source'] = str(data_source.get('path') or '')
    response['X-Generated-At'] = timezone.now().isoformat()
    return response


def render_csv(built):
    """Render a built report as a ``text/csv`` attachment."""
    rows = enforce_row_cap(built)
    buffer = io.StringIO(newline='')
    writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL, lineterminator='\r\n')
    writer.writerow(_row(built.table['columns']))
    for row in rows:
        writer.writerow(_row(row))
    # utf-8-sig so Excel on Windows reads the file as UTF-8 rather than the
    # local codepage, which mangles the emoji post bodies are full of.
    body = buffer.getvalue().encode('utf-8-sig')
    return _attach(HttpResponse(body, content_type=CSV_MIME), built, 'csv', len(rows))


def render_xlsx(built):
    """Render a built report as ``.xlsx`` in openpyxl's write-only mode.

    Raises :class:`XlsxUnavailable` when openpyxl cannot be imported, so the
    caller answers 503 naming the package instead of leaking an ImportError.
    """
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise XlsxUnavailable(str(exc)) from exc

    rows = enforce_row_cap(built)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=built.key[:31] or 'report')
    sheet.append(_row(built.table['columns']))
    for row in rows:
        sheet.append(_row(row))

    # Second sheet: where the numbers came from. A downloaded file outlives the
    # response headers, so provenance has to live inside it too or it is lost.
    provenance = workbook.create_sheet(title='provenance')
    provenance.append(['key', 'value'])
    window = built.payload.get('window') or {}
    data_source = built.payload.get('data_source') or {}
    for label, value in (
        ('report', built.key),
        ('label', built.label),
        ('from', window.get('from')),
        ('to', window.get('to')),
        ('days', window.get('days')),
        ('rows', len(rows)),
        ('data_source_path', data_source.get('path')),
        ('data_source_note', data_source.get('note')),
        ('generated_at', timezone.now().isoformat()),
        ('query_params', ' '.join(f'{k}={v}' for k, v in sorted(built.params.items()))),
    ):
        provenance.append([label, value])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return _attach(
        HttpResponse(buffer.getvalue(), content_type=XLSX_MIME),
        built, 'xlsx', len(rows),
    )


def csv_response(report_key, params):
    """Build + render in one step. Convenience for non-view callers."""
    return render_csv(build_report(report_key, params))


def xlsx_response(report_key, params):
    """Build + render in one step. Raises :class:`XlsxUnavailable`."""
    return render_xlsx(build_report(report_key, params))


class XlsxUnavailable(Exception):
    """openpyxl is not installed. Rendered as a 503 that names the package."""

    def __init__(self, detail=''):
        super().__init__(detail)
        self.detail = detail


XLSX_MISSING_MESSAGE = (
    'xlsx export is unavailable on this deployment: the "openpyxl" package is '
    'not installed. It is declared in requirements/base.txt; rebuild the image '
    'or reinstall requirements. CSV export of the same report is unaffected.'
)